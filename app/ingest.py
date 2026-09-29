from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import DocumentChunk, SourceSummary


_TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _extract_csv(data: str) -> str:
    rows = csv.reader(io.StringIO(data))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _extract_docx(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text(path: Path) -> str:
    path = Path(path)
    extension = path.suffix.lower()
    if extension in _TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8")
    elif extension == ".csv":
        text = _extract_csv(path.read_text(encoding="utf-8-sig"))
    elif extension == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        text = json.dumps(value, ensure_ascii=False, indent=2)
    elif extension == ".pdf":
        text = _extract_pdf(path)
    elif extension == ".docx":
        text = _extract_docx(path)
    else:
        raise ValueError(f"Unsupported file extension: {extension or '<none>'}")
    return _normalize_text(text)


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size")
    text = _normalize_text(text)
    if not text:
        return []
    step = chunk_size - overlap
    return [
        text[start : start + chunk_size]
        for start in range(0, len(text), step)
        if start == 0 or len(text[start : start + chunk_size]) > overlap
    ]


def _source_id(path: Path, content: bytes) -> str:
    digest = hashlib.sha256(str(path.absolute()).encode("utf-8") + b"\0" + content).hexdigest()
    return f"source-{digest[:16]}"


def ingest_file(
    path: Path, chunk_size: int = 900, overlap: int = 120
) -> tuple[SourceSummary, list[DocumentChunk]]:
    path = Path(path)
    content = path.read_bytes()
    source_id = _source_id(path, content)
    chunks = chunk_text(extract_text(path), chunk_size=chunk_size, overlap=overlap)
    source_name = path.name
    documents = [
        DocumentChunk(
            chunk_id=f"{source_id}-{index}",
            source_id=source_id,
            source_name=source_name,
            chunk_index=index,
            text=chunk,
        )
        for index, chunk in enumerate(chunks)
    ]
    summary = SourceSummary(
        source_id=source_id,
        name=source_name,
        chunk_count=len(documents),
        file_type=path.suffix.lower(),
        path=str(path),
        size_bytes=len(content),
    )
    return summary, documents
